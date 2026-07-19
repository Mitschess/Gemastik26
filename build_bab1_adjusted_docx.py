from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "BAB1_SatuKelas_Disesuaikan_Dengan_Gambar.docx"


def set_font(run, name="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.333
    r = p.add_run(text)
    set_font(r, size=11)
    return p


def add_placeholder(doc, no, title, desc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"[Gambar {no}. {title}]")
    set_font(r, size=10, bold=True, color="1F4D78")

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(10)
    r2 = p2.add_run(desc)
    set_font(r2, size=9, italic=True, color="555555")


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(item)
        set_font(r, size=11)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.333

for style_name, size, color, before, after in [
    ("Heading 1", 16, "2E74B5", 18, 10),
    ("Heading 2", 13, "2E74B5", 12, 6),
    ("Heading 3", 12, "1F4D78", 8, 4),
]:
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("SatuKelas")
set_font(r, size=22, bold=True, color="0B2545")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(16)
r = subtitle.add_run("Ekosistem Pembelajaran Offline Berbasis Edge AI dan LoRa-GPS untuk Daerah 3T dan Kondisi Bencana")
set_font(r, size=12, italic=True, color="1F4D78")

doc.add_heading("BAB I", level=1)
doc.add_heading("PENDAHULUAN", level=1)

add_body(doc, "Transformasi digital pendidikan telah menjadi kebutuhan mendesak dalam meningkatkan kualitas, pemerataan, dan ketahanan layanan pembelajaran. Berbagai Learning Management System (LMS) seperti Google Classroom dan Moodle telah membantu guru dalam mendistribusikan materi, mengelola tugas, menyelenggarakan kuis, mencatat presensi, serta memantau capaian belajar peserta didik. Namun, pemanfaatan LMS pada umumnya masih sangat bergantung pada ketersediaan internet yang stabil, perangkat pendukung, dan kemampuan sekolah dalam mengelola infrastruktur digital (Dhawan, 2020).")
add_body(doc, "Kondisi tersebut menimbulkan tantangan besar bagi sekolah di daerah terdepan, terluar, dan tertinggal (3T). Berdasarkan Peraturan Presiden Republik Indonesia Nomor 63 Tahun 2020, daerah tertinggal merupakan wilayah yang masih menghadapi keterbatasan dalam aspek pembangunan, termasuk infrastruktur, aksesibilitas, dan pelayanan publik. Dalam konteks pendidikan, keterbatasan tersebut tampak pada akses internet yang tidak merata, biaya konektivitas yang relatif tinggi, keterbatasan tenaga pendidik, serta minimnya dukungan teknis untuk menjalankan pembelajaran digital secara berkelanjutan (UNICEF Indonesia, 2021).")
add_placeholder(doc, 1, "Konteks Daerah 3T dan Kesenjangan Konektivitas", "Sisipkan peta sebaran daerah 3T atau infografik akses internet pendidikan di Indonesia.")

add_body(doc, "Di sisi lain, pemerintah telah mendorong digitalisasi sekolah melalui penyediaan perangkat teknologi, termasuk Chromebook dan perangkat pendukung pembelajaran lainnya. Akan tetapi, perangkat tersebut belum selalu dapat dimanfaatkan secara optimal apabila sekolah berada pada wilayah dengan konektivitas rendah. Akibatnya, teknologi yang seharusnya memperluas akses belajar justru berisiko menjadi kurang efektif karena sistem pembelajaran digital yang tersedia tidak dirancang untuk kondisi offline atau jaringan terbatas (Astari & Yulianto, 2025).")
add_body(doc, "Permasalahan serupa juga muncul pada situasi bencana alam. Ketika gempa bumi, banjir, erupsi gunung api, longsor, atau keadaan darurat lain terjadi, akses internet dan listrik sering kali terganggu. Sekolah, pos pengungsian, maupun pusat layanan masyarakat tetap membutuhkan media pembelajaran, distribusi informasi, dan pencatatan aktivitas pendidikan yang dapat berjalan secara lokal. Oleh karena itu, diperlukan sistem pembelajaran yang tidak hanya mendukung daerah 3T, tetapi juga mampu berfungsi sebagai infrastruktur pendidikan darurat saat jaringan internet publik tidak tersedia (Dhawan, 2020).")
add_body(doc, "Selain pemerataan akses, pendidikan saat ini juga membutuhkan pendekatan yang mendukung literasi digital dan pembelajaran coding. Kemampuan berpikir komputasional, pemrograman dasar, dan pemecahan masalah berbasis teknologi semakin penting bagi peserta didik. Namun, pembelajaran coding sering kali terkendala kebutuhan akses internet, platform daring, dan sumber belajar interaktif. Sistem pembelajaran lokal yang dapat menyediakan materi, latihan, evaluasi, serta asisten belajar berbasis kecerdasan buatan akan membantu siswa mempelajari coding meskipun berada pada lingkungan dengan koneksi terbatas (Almaiah et al., 2020).")

add_body(doc, "Berangkat dari permasalahan tersebut, dikembangkan SatuKelas, yaitu LMS berbasis jaringan WiFi lokal yang dijalankan melalui node komputasi seperti Raspberry Pi. SatuKelas memungkinkan siswa dan guru mengakses materi, tugas, kuis, presensi, forum, serta penilaian melalui jaringan lokal sekolah tanpa bergantung pada internet publik. SatuKelas juga dirancang untuk memanfaatkan kecerdasan buatan, seperti model Qwen, OCR, dan Retrieval-Augmented Generation (RAG), sehingga sistem tidak hanya menyimpan materi, tetapi juga dapat membantu guru menyusun soal, merangkum materi, membaca dokumen hasil pindai, dan menjawab pertanyaan siswa berdasarkan materi yang diunggah guru. Pendekatan tersebut sejalan dengan konsep edge intelligence, yaitu pemrosesan data dan kecerdasan buatan langsung pada perangkat lokal sehingga mampu mengurangi ketergantungan terhadap layanan cloud dan meningkatkan efisiensi pada lingkungan dengan keterbatasan konektivitas (Zhou et al., 2019).")
add_placeholder(doc, 2, "Arsitektur SatuKelas Berbasis Raspberry Pi", "Sisipkan diagram: Raspberry Pi sebagai server lokal, guru dan siswa mengakses LMS melalui WiFi lokal.")

add_body(doc, "Untuk mendukung pemantauan distribusi dan pemanfaatan perangkat, SatuKelas juga dapat diintegrasikan dengan GPS dan LoRa. GPS digunakan untuk mencatat lokasi penggunaan node, sedangkan LoRa digunakan untuk mengirimkan data ringan seperti identitas node, status perangkat, lokasi, durasi aktif, dan ringkasan penggunaan. Dengan pendekatan ini, pemerintah, dinas pendidikan, atau lembaga pengelola program tidak harus menempatkan Raspberry Pi di sisi penerima. Pihak pemantau cukup menggunakan gateway LoRa atau LoRaWAN receiver yang terhubung ke server atau dashboard untuk menerima telemetri dari node SatuKelas. Raspberry Pi tetap berperan sebagai server lokal di sekolah, sedangkan LoRa berperan sebagai kanal komunikasi jarak jauh berdaya rendah untuk monitoring (Hong & Varghese, 2019).")
add_placeholder(doc, 3, "Skema Monitoring LoRa-GPS", "Sisipkan diagram: node SatuKelas + GPS mengirim metadata melalui LoRa ke gateway dan dashboard pemantauan.")

add_body(doc, "Namun, LoRa memiliki keterbatasan bandwidth sehingga tidak ideal untuk mengirim file besar seperti video, PDF berukuran besar, dataset, atau model AI. LoRa lebih tepat digunakan untuk komunikasi metadata, status, lokasi, dan permintaan sinkronisasi. Sementara itu, transfer file pembelajaran tetap lebih sesuai dilakukan melalui WiFi lokal, penyimpanan eksternal, sinkronisasi berkala saat internet tersedia, atau mekanisme distribusi offline lainnya. Kombinasi LMS lokal, AI, OCR, RAG, GPS, dan LoRa menjadikan SatuKelas sebagai solusi pembelajaran digital yang inklusif, adaptif, dan relevan untuk daerah 3T, kondisi bencana, serta penguatan pembelajaran teknologi (Hong & Varghese, 2019).")

doc.add_heading("1.2 Rumusan Masalah", level=2)
add_body(doc, "Adapun rumusan masalah yang diangkat dalam penelitian ini adalah:")
add_numbered(doc, [
    "Bagaimana merancang LMS yang dapat berjalan secara lokal tanpa ketergantungan pada internet publik untuk mendukung pembelajaran di daerah 3T dan situasi bencana alam?",
    "Bagaimana mengintegrasikan kecerdasan buatan ke dalam SatuKelas agar dapat membantu guru membuat soal, memahami materi yang diunggah, melakukan OCR, serta memberikan bantuan belajar kepada siswa?",
    "Bagaimana SatuKelas dapat mendukung pembelajaran coding dan literasi digital pada lingkungan dengan keterbatasan konektivitas?",
    "Bagaimana mengintegrasikan GPS dan LoRa sebagai sistem monitoring penggunaan node SatuKelas sehingga pemerintah atau dinas pendidikan dapat memantau lokasi, status, dan pemanfaatan perangkat secara efisien tanpa bergantung pada internet?",
])

doc.add_heading("1.3 Tujuan", level=2)
add_body(doc, "Adapun tujuan yang ingin dicapai dalam pengembangan SatuKelas adalah:")
add_numbered(doc, [
    "Mengembangkan SatuKelas sebagai LMS lokal yang dapat digunakan tanpa internet publik untuk mendukung proses pembelajaran di daerah 3T dan wilayah terdampak bencana.",
    "Menerapkan kecerdasan buatan pada SatuKelas untuk membantu pembuatan soal, penyusunan ringkasan, OCR materi, pencarian jawaban berbasis dokumen guru, dan pendampingan belajar siswa.",
    "Menyediakan lingkungan pembelajaran coding yang dapat diakses secara lokal, sehingga siswa tetap dapat mempelajari pemrograman dan literasi digital meskipun konektivitas terbatas.",
    "Mengembangkan sistem monitoring node SatuKelas menggunakan GPS dan LoRa untuk mendukung evaluasi pemanfaatan perangkat oleh pemerintah maupun sekolah.",
])

doc.add_heading("1.4 Urgensi", level=2)
add_body(doc, "Pengembangan SatuKelas menjadi penting karena pemerataan pendidikan digital tidak dapat hanya mengandalkan model pembelajaran daring berbasis internet. Daerah 3T dan wilayah terdampak bencana membutuhkan sistem yang tetap dapat berjalan ketika koneksi publik tidak tersedia. Tanpa solusi offline-first, siswa di wilayah tersebut berpotensi semakin tertinggal dalam akses materi, evaluasi, dan pengalaman belajar digital.")
add_body(doc, "Urgensi lain terletak pada pemanfaatan teknologi yang sudah tersedia. Perangkat seperti Chromebook, laptop sekolah, atau gawai siswa dapat menjadi lebih bernilai apabila tersedia server lokal yang mampu menyediakan layanan pembelajaran digital. SatuKelas memungkinkan perangkat tersebut digunakan dalam jaringan WiFi lokal untuk mengakses materi, mengerjakan tugas, mengikuti kuis, dan berlatih coding tanpa internet.")
add_placeholder(doc, 4, "Pemanfaatan Perangkat Sekolah dalam Jaringan Lokal", "Sisipkan foto/mockup Chromebook, laptop, atau gawai siswa yang terhubung ke server lokal SatuKelas.")

add_body(doc, "Integrasi AI juga penting untuk mengurangi beban guru dan meningkatkan kualitas pembelajaran. Dengan OCR dan RAG, materi berupa PDF, dokumen, atau hasil pindai dapat diubah menjadi sumber pengetahuan yang dapat ditanyakan kembali oleh siswa. Dengan generator soal dan asisten belajar, guru memperoleh dukungan dalam menyusun evaluasi dan siswa memperoleh bantuan belajar yang lebih personal.")
add_placeholder(doc, 5, "Alur AI Lokal untuk Pembelajaran", "Sisipkan diagram: dokumen guru -> OCR/RAG -> model AI lokal -> ringkasan, soal, dan jawaban berbasis materi.")

add_body(doc, "Dari sisi tata kelola, integrasi LoRa dan GPS memberi nilai strategis karena pengelola program dapat mengetahui apakah node benar-benar aktif, berapa lama digunakan, dan di lokasi mana perangkat dimanfaatkan. Data ringan tersebut dapat membantu pemerintah mengevaluasi distribusi perangkat, efektivitas program pengiriman guru atau perangkat belajar, serta kebutuhan dukungan lanjutan di lapangan.")
add_body(doc, "Beberapa solusi LMS offline seperti MoodleBox, Kolibri, RACHEL, dan LibraryBox telah tersedia untuk mendukung pembelajaran tanpa internet. Namun, sebagian besar masih berfokus pada distribusi materi pembelajaran dan belum mengintegrasikan kecerdasan buatan lokal, OCR, monitoring perangkat berbasis LoRa dan GPS, maupun dukungan pembelajaran coding dalam satu ekosistem. Oleh karena itu, masih terdapat kebutuhan akan platform pembelajaran offline yang lebih adaptif terhadap kondisi daerah 3T dan situasi bencana.")

doc.add_heading("1.5 Manfaat", level=2)
add_numbered(doc, [
    "Menyediakan platform Learning Management System (LMS) berbasis jaringan lokal yang memungkinkan proses pembelajaran digital tetap berlangsung tanpa bergantung pada internet publik, sehingga mendukung pemerataan pendidikan di daerah 3T maupun wilayah terdampak bencana.",
    "Meningkatkan efektivitas proses belajar mengajar melalui integrasi teknologi Artificial Intelligence (AI), Optical Character Recognition (OCR), dan Retrieval-Augmented Generation (RAG) untuk membantu guru menyusun materi, soal, serta mendampingi siswa dalam memahami materi pembelajaran.",
    "Mendukung pembelajaran coding dan literasi digital melalui lingkungan belajar berbasis offline-first yang dapat diakses menggunakan perangkat seperti Chromebook, laptop, maupun gawai dalam jaringan WiFi lokal.",
    "Mempermudah pemerintah, dinas pendidikan, maupun pengelola program dalam memantau lokasi, status, durasi aktif, dan tingkat pemanfaatan perangkat melalui integrasi GPS dan LoRa sebagai dasar evaluasi pemerataan pendidikan digital.",
    "Menghasilkan prototipe LMS berbasis edge computing yang mengintegrasikan AI, GPS, dan LoRa sebagai solusi pembelajaran digital yang inklusif, adaptif, dan berkelanjutan bagi wilayah dengan keterbatasan infrastruktur.",
])

doc.add_page_break()
doc.add_heading("DAFTAR GAMBAR YANG DISARANKAN", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
for row in table.rows:
    for cell, width in zip(row.cells, [1.0, 2.8, 2.7]):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)

headers = ["No.", "Gambar", "Letak dan Fungsi"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = text
    shade_cell(cell, "F4F6F9")
    for p in cell.paragraphs:
        for run in p.runs:
            set_font(run, size=10, bold=True, color="0B2545")

image_rows = [
    ("1", "Peta daerah 3T atau infografik kesenjangan akses internet.", "Setelah pembahasan masalah daerah 3T untuk memperkuat latar belakang."),
    ("2", "Diagram arsitektur Raspberry Pi sebagai server lokal SatuKelas.", "Setelah penjelasan solusi agar pembaca memahami hubungan server, WiFi, guru, dan siswa."),
    ("3", "Skema LoRa-GPS untuk monitoring node.", "Setelah paragraf integrasi LoRa dan GPS untuk menjelaskan alur telemetri."),
    ("4", "Foto/mockup perangkat sekolah dalam jaringan lokal.", "Pada bagian urgensi pemanfaatan Chromebook, laptop, atau gawai siswa."),
    ("5", "Diagram AI lokal: OCR, RAG, dan model Qwen.", "Pada bagian urgensi AI agar fitur bantuan guru dan siswa lebih mudah dipahami."),
]
for row in image_rows:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].text = text
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[i])
        for p in cells[i].paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_font(run, size=10)

doc.add_heading("DAFTAR PUSTAKA", level=1)
refs = [
    "Almaiah, M. A., Al-Khasawneh, A., & Althunibat, A. (2020). Exploring the critical challenges and factors influencing the e-learning system usage during the COVID-19 pandemic. Education and Information Technologies, 25(6), 5261-5280. https://doi.org/10.1007/s10639-020-10219-y",
    "Astari, & Yulianto, D. (2025). Bridging the digital divide in education: Disparities in Google Classroom utilization and technical challenges among urban and rural teachers. Journal of Education Technology, 9(2), 258-270. https://doi.org/10.23887/jet.v9i2.92897",
    "Dhawan, S. (2020). Online learning: A panacea in the time of COVID-19 crisis. Journal of Educational Technology Systems, 49(1), 5-22. https://doi.org/10.1177/0047239520934018",
    "Hong, C. H., & Varghese, B. (2019). Resource management in fog/edge computing: A survey. ACM Computing Surveys, 52(5), 1-37. https://doi.org/10.1145/3326066",
    "Khalil, H., Ebner, M., & Taraghi, B. (2020). Development and evaluation of open source learning management systems. Education and Information Technologies, 25, 1-20.",
    "Learning Equality. (n.d.). Kolibri. https://learningequality.org/kolibri/",
    "MoodleBox. (n.d.). MoodleBox. https://moodlebox.net/en/",
    "UNICEF Indonesia. (2021). Situation analysis on digital learning in Indonesia. UNICEF Indonesia. https://www.unicef.org/indonesia/reports/situation-analysis-digital-learning-indonesia",
    "Zhou, Z., Chen, X., Li, E., Zeng, L., Luo, K., Zhang, J. (2019). Edge intelligence: Paving the last mile of artificial intelligence with edge computing. Proceedings of the IEEE, 107(8), 1738-1762. https://doi.org/10.1109/JPROC.2019.2918951",
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(ref)
    set_font(r, size=10)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("SatuKelas - BAB I Disesuaikan")
set_font(fr, size=9, color="555555")

doc.save(OUT)
print(OUT)
