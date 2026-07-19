from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "Esai_SatuKelas_LoRa_RaspberryPi_GPS.docx"


def set_font(run, name="Calibri", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_para(doc, text, style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.333
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_image_placeholder(doc, number, title, suggestion):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"[Gambar {number}: {title}]")
    set_font(r, size=10, bold=True, color="1F4D78")

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(10)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"Sisipkan: {suggestion}")
    set_font(r2, size=9, italic=True, color="555555")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.333

for style_name, size, color, before, after in [
    ("Heading 1", 16, "2E74B5", 18, 10),
    ("Heading 2", 13, "2E74B5", 12, 6),
    ("Heading 3", 12, "1F4D78", 8, 4),
]:
    st = styles[style_name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(6)
r = title.add_run("SatuKelas")
set_font(r, size=22, bold=True, color="0B2545")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
r = subtitle.add_run(
    "Ekosistem Pembelajaran Offline Berbasis Edge AI, LoRa, Raspberry Pi, dan GPS untuk Daerah 3T dan Kondisi Bencana"
)
set_font(r, size=12, italic=True, color="1F4D78")

add_para(
    doc,
    "Transformasi digital pendidikan telah menjadi kebutuhan penting dalam meningkatkan kualitas, pemerataan, dan ketahanan layanan pembelajaran. Berbagai Learning Management System (LMS) seperti Google Classroom dan Moodle telah membantu guru membagikan materi, mengelola tugas, menyelenggarakan kuis, mencatat presensi, serta memantau perkembangan belajar peserta didik. Namun, sebagian besar platform tersebut masih bergantung pada ketersediaan internet yang stabil. Kondisi ini menjadi tantangan besar bagi sekolah di daerah terdepan, terluar, dan tertinggal (3T), karena tidak semua wilayah memiliki akses internet yang memadai."
)

add_image_placeholder(
    doc,
    1,
    "Kesenjangan Akses Internet Pendidikan",
    "peta daerah 3T Indonesia atau infografik kesenjangan akses internet sekolah.",
)

add_para(
    doc,
    "Kesenjangan akses internet menyebabkan digitalisasi pendidikan belum dapat dirasakan secara merata. Di satu sisi, pemerintah telah berupaya menyediakan perangkat teknologi seperti Chromebook dan perangkat pendukung pembelajaran lainnya. Akan tetapi, perangkat tersebut belum dapat dimanfaatkan secara optimal apabila sekolah berada di wilayah dengan konektivitas rendah. Akibatnya, teknologi yang seharusnya menjadi alat pemerataan pendidikan justru berisiko menciptakan ketimpangan baru antara sekolah yang memiliki akses internet baik dan sekolah yang masih terbatas secara infrastruktur."
)
add_para(
    doc,
    "Permasalahan serupa juga muncul dalam kondisi bencana alam. Ketika terjadi gempa bumi, banjir, longsor, erupsi gunung api, atau keadaan darurat lain, jaringan internet dan listrik sering kali terganggu. Padahal, kegiatan pendidikan tetap perlu berjalan, baik di sekolah, pos pengungsian, maupun pusat layanan masyarakat. Oleh karena itu, dibutuhkan sistem pembelajaran yang mampu berjalan secara mandiri tanpa bergantung sepenuhnya pada internet publik."
)

doc.add_heading("Gagasan Solusi", level=1)
add_para(
    doc,
    "SatuKelas hadir sebagai solusi pembelajaran digital berbasis jaringan lokal. Sistem ini dirancang sebagai LMS offline-first yang dapat dijalankan melalui perangkat komputasi kecil seperti Raspberry Pi. Raspberry Pi berperan sebagai server lokal yang menyimpan materi, tugas, kuis, presensi, forum, penilaian, serta fitur pembelajaran lain. Siswa dan guru dapat mengakses SatuKelas melalui jaringan WiFi lokal menggunakan Chromebook, laptop, tablet, atau gawai tanpa perlu terhubung ke internet."
)

add_image_placeholder(
    doc,
    2,
    "Arsitektur SatuKelas Berbasis Raspberry Pi",
    "diagram sederhana: Raspberry Pi sebagai server lokal, guru dan siswa terhubung lewat WiFi lokal.",
)

add_para(
    doc,
    "Keunggulan utama SatuKelas terletak pada kemampuannya menghadirkan pengalaman pembelajaran digital di lingkungan dengan konektivitas terbatas. Guru dapat mengunggah materi, membuat soal, memberikan tugas, dan memantau hasil belajar siswa secara lokal. Siswa juga dapat membaca materi, mengerjakan kuis, mengumpulkan tugas, serta mengikuti pembelajaran coding tanpa bergantung pada platform daring. Dengan demikian, perangkat yang sudah tersedia di sekolah dapat digunakan secara lebih optimal."
)

doc.add_heading("Integrasi Edge AI dalam Pembelajaran Lokal", level=1)
add_para(
    doc,
    "Selain sebagai LMS lokal, SatuKelas juga dapat diperkuat dengan teknologi kecerdasan buatan berbasis edge AI. Melalui integrasi model AI lokal, OCR, dan Retrieval-Augmented Generation (RAG), sistem ini dapat membantu guru menyusun soal, merangkum materi, membaca dokumen hasil pindai, serta menjawab pertanyaan siswa berdasarkan materi yang telah diunggah. Karena pemrosesan dilakukan secara lokal, sekolah tidak selalu membutuhkan layanan cloud atau internet untuk menjalankan fitur bantuan belajar tersebut."
)

add_image_placeholder(
    doc,
    3,
    "Alur Kerja AI Lokal",
    "diagram alur: materi guru -> OCR/RAG -> model AI lokal -> jawaban siswa atau bantuan pembuatan soal.",
)

doc.add_heading("Peran LoRa, GPS, dan Raspberry Pi", level=1)
add_para(
    doc,
    "Integrasi Raspberry Pi, LoRa, dan GPS menjadi nilai tambah penting dalam pengembangan SatuKelas. Raspberry Pi digunakan sebagai pusat layanan pembelajaran lokal di sekolah. GPS berfungsi mencatat lokasi penggunaan node SatuKelas, sedangkan LoRa digunakan untuk mengirim data ringan seperti identitas perangkat, status aktif, lokasi, durasi penggunaan, dan ringkasan pemanfaatan. Data ini dapat diterima oleh gateway LoRa atau LoRaWAN receiver yang terhubung ke dashboard pemantauan."
)

add_image_placeholder(
    doc,
    4,
    "Monitoring Node dengan LoRa dan GPS",
    "diagram: node SatuKelas + GPS mengirim metadata melalui LoRa ke gateway dan dashboard dinas pendidikan.",
)

add_para(
    doc,
    "Dengan sistem tersebut, pemerintah, dinas pendidikan, atau pengelola program dapat mengetahui apakah perangkat benar-benar digunakan, di mana perangkat berada, dan seberapa sering sistem dimanfaatkan. Hal ini penting untuk mengevaluasi efektivitas distribusi perangkat pendidikan digital, terutama di daerah 3T dan wilayah terdampak bencana. LoRa tidak digunakan untuk mengirim file besar seperti video atau modul pembelajaran, tetapi lebih tepat untuk mengirim metadata dan informasi pemantauan karena memiliki konsumsi daya rendah serta jangkauan komunikasi yang luas."
)

doc.add_heading("Dampak bagi Literasi Digital dan Pembelajaran Coding", level=1)
add_para(
    doc,
    "SatuKelas juga relevan untuk mendukung pembelajaran coding dan literasi digital. Di banyak wilayah, pembelajaran pemrograman sering terkendala akses internet dan minimnya platform interaktif. Dengan SatuKelas, materi coding, latihan, kuis, dan asisten belajar dapat diakses secara lokal. Siswa tetap dapat belajar berpikir komputasional dan pemrograman dasar meskipun berada di lingkungan dengan keterbatasan konektivitas."
)

add_image_placeholder(
    doc,
    5,
    "Aktivitas Belajar Coding Offline",
    "foto atau mockup siswa menggunakan Chromebook/laptop untuk latihan coding melalui jaringan lokal.",
)

doc.add_heading("Kesimpulan", level=1)
add_para(
    doc,
    "Melalui kombinasi LMS lokal, edge AI, OCR, RAG, Raspberry Pi, GPS, dan LoRa, SatuKelas menawarkan model transformasi digital pendidikan yang lebih inklusif dan kontekstual. Sistem ini tidak hanya menjawab persoalan keterbatasan internet, tetapi juga membantu optimalisasi perangkat sekolah, mendukung pembelajaran darurat saat bencana, serta menyediakan mekanisme pemantauan yang efisien bagi pemerintah."
)
add_para(
    doc,
    "Dengan demikian, SatuKelas dapat menjadi inovasi pendidikan digital yang sesuai dengan kebutuhan Indonesia, khususnya bagi daerah 3T dan wilayah dengan infrastruktur terbatas. Transformasi digital pendidikan tidak seharusnya hanya bergantung pada internet, tetapi perlu dirancang agar mampu menyesuaikan diri dengan kondisi nyata masyarakat. SatuKelas membuktikan bahwa pembelajaran digital tetap dapat berjalan secara mandiri, adaptif, dan berkelanjutan melalui pemanfaatan teknologi lokal yang tepat guna."
)

doc.add_page_break()
doc.add_heading("Daftar Gambar yang Disarankan", level=1)

items = [
    ("Gambar 1", "Peta daerah 3T atau infografik kesenjangan internet", "Diletakkan setelah paragraf pembuka agar pembaca langsung memahami masalah utama."),
    ("Gambar 2", "Diagram arsitektur Raspberry Pi sebagai server lokal", "Diletakkan pada bagian gagasan solusi untuk menjelaskan cara kerja SatuKelas."),
    ("Gambar 3", "Diagram alur AI lokal: OCR, RAG, dan model AI", "Diletakkan pada bagian edge AI agar fitur AI lebih mudah dipahami."),
    ("Gambar 4", "Diagram LoRa-GPS untuk monitoring node", "Diletakkan pada bagian LoRa, GPS, dan Raspberry Pi."),
    ("Gambar 5", "Foto/mockup siswa belajar coding offline", "Diletakkan pada bagian literasi digital dan pembelajaran coding."),
]

table = doc.add_table(rows=1, cols=3)
set_table_width(table, [1.1, 2.7, 2.7])
hdr = table.rows[0].cells
for idx, text in enumerate(["Posisi", "Gambar yang Dibutuhkan", "Alasan Penempatan"]):
    hdr[idx].text = text
    set_cell_shading(hdr[idx], "F4F6F9")
    for p in hdr[idx].paragraphs:
        for run in p.runs:
            set_font(run, bold=True, color="0B2545")

for row in items:
    cells = table.add_row().cells
    for idx, text in enumerate(row):
        cells[idx].text = text
        for p in cells[idx].paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_font(run, size=10)

doc.add_heading("Catatan Penyisipan Gambar", level=2)
for note in [
    "Gunakan gambar dengan resolusi cukup tinggi agar tidak pecah saat dicetak.",
    "Untuk diagram, gunakan warna sederhana dan beri label yang jelas.",
    "Jika memakai foto nyata, pastikan tidak menampilkan wajah siswa tanpa izin.",
    "Gambar LoRa, Raspberry Pi, GPS, dan dashboard sebaiknya berupa diagram agar mudah dinilai sebagai rancangan sistem.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(note)
    set_font(r)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("SatuKelas - Esai Inovasi Pendidikan Digital")
set_font(fr, size=9, color="555555")

doc.save(OUT)
print(OUT)
